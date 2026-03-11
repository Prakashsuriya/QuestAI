import os
import tempfile
from datetime import datetime
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
import pandas as pd
from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors
from app.models import Question, Answer, Citation

class ExportService:
    """Service for exporting questionnaires in various formats."""
    
    def export_to_docx(self, questionnaire) -> str:
        """Export questionnaire to Word document."""
        doc = Document()
        
        # Title
        title = doc.add_heading('Completed Questionnaire', 0)
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        
        # Metadata
        doc.add_paragraph(f"Original File: {questionnaire.original_filename}")
        doc.add_paragraph(f"Export Date: {datetime.now().strftime('%Y-%m-%d %H:%M')}")
        doc.add_paragraph(f"Status: {questionnaire.status}")
        doc.add_paragraph()
        
        # Questions and Answers
        questions = Question.query.filter_by(questionnaire_id=questionnaire.id).order_by(Question.question_number).all()
        
        for question in questions:
            # Question
            q_para = doc.add_paragraph()
            q_run = q_para.add_run(f"Q{question.question_number}: ")
            q_run.bold = True
            q_para.add_run(question.question_text)
            
            # Answer
            if question.answers:
                answer = question.answers[0]
                a_para = doc.add_paragraph()
                a_run = a_para.add_run("Answer: ")
                a_run.bold = True
                a_para.add_run(answer.answer_text)
                
                # Citations
                if answer.citations:
                    c_para = doc.add_paragraph()
                    c_run = c_para.add_run("Citations: ")
                    c_run.italic = True
                    c_run.font.size = Pt(10)
                    citation_texts = []
                    for citation in answer.citations:
                        cite_str = f"{citation.document.original_filename}"
                        if citation.page_number:
                            cite_str += f", p.{citation.page_number}"
                        citation_texts.append(cite_str)
                    c_para.add_run("; ".join(citation_texts))
                    c_para.paragraph_format.left_indent = Inches(0.5)
                
                # Confidence
                if answer.confidence_score is not None:
                    conf_para = doc.add_paragraph()
                    conf_run = conf_para.add_run(f"Confidence: {answer.confidence_score:.0%}")
                    conf_run.font.size = Pt(9)
                    conf_run.font.color.rgb = None  # Will use default color
                    conf_para.paragraph_format.left_indent = Inches(0.5)
                
                # Edited indicator
                if answer.is_edited:
                    edit_para = doc.add_paragraph()
                    edit_run = edit_para.add_run("(Edited by user)")
                    edit_run.italic = True
                    edit_run.font.size = Pt(9)
                    edit_para.paragraph_format.left_indent = Inches(0.5)
            else:
                a_para = doc.add_paragraph()
                a_run = a_para.add_run("Answer: ")
                a_run.bold = True
                a_para.add_run("Not answered")
            
            doc.add_paragraph()  # Spacing between Q&A pairs
        
        # Save to temp file
        output_path = os.path.join(tempfile.gettempdir(), f"questionnaire_{questionnaire.id}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.docx")
        doc.save(output_path)
        
        return output_path
    
    def export_to_excel(self, questionnaire) -> str:
        """Export questionnaire to Excel."""
        questions = Question.query.filter_by(questionnaire_id=questionnaire.id).order_by(Question.question_number).all()
        
        data = []
        for question in questions:
            row = {
                'Question Number': question.question_number,
                'Category': question.category,
                'Question': question.question_text,
                'Answer': '',
                'Citations': '',
                'Confidence': '',
                'Edited': 'No'
            }
            
            if question.answers:
                answer = question.answers[0]
                row['Answer'] = answer.answer_text
                row['Confidence'] = f"{answer.confidence_score:.0%}" if answer.confidence_score else "N/A"
                row['Edited'] = "Yes" if answer.is_edited else "No"
                
                if answer.citations:
                    citation_texts = []
                    for citation in answer.citations:
                        cite_str = f"{citation.document.original_filename}"
                        if citation.page_number:
                            cite_str += f", p.{citation.page_number}"
                        citation_texts.append(cite_str)
                    row['Citations'] = "; ".join(citation_texts)
            else:
                row['Answer'] = "Not answered"
            
            data.append(row)
        
        df = pd.DataFrame(data)
        
        output_path = os.path.join(tempfile.gettempdir(), f"questionnaire_{questionnaire.id}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx")
        df.to_excel(output_path, index=False, sheet_name='Questionnaire')
        
        return output_path
    
    def export_to_pdf(self, questionnaire) -> str:
        """Export questionnaire to PDF."""
        output_path = os.path.join(tempfile.gettempdir(), f"questionnaire_{questionnaire.id}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.pdf")
        
        doc = SimpleDocTemplate(output_path, pagesize=letter,
                               rightMargin=72, leftMargin=72,
                               topMargin=72, bottomMargin=18)
        
        styles = getSampleStyleSheet()
        story = []
        
        # Custom styles
        title_style = ParagraphStyle(
            'CustomTitle',
            parent=styles['Heading1'],
            fontSize=18,
            textColor=colors.HexColor('#1a1a1a'),
            spaceAfter=30,
            alignment=1  # Center
        )
        
        question_style = ParagraphStyle(
            'QuestionStyle',
            parent=styles['Normal'],
            fontSize=11,
            textColor=colors.HexColor('#333333'),
            spaceAfter=6,
            spaceBefore=12
        )
        
        answer_style = ParagraphStyle(
            'AnswerStyle',
            parent=styles['Normal'],
            fontSize=10,
            textColor=colors.HexColor('#000000'),
            leftIndent=20,
            spaceAfter=6
        )
        
        citation_style = ParagraphStyle(
            'CitationStyle',
            parent=styles['Normal'],
            fontSize=9,
            textColor=colors.HexColor('#666666'),
            leftIndent=20,
            spaceAfter=12
        )
        
        # Title
        story.append(Paragraph("Completed Questionnaire", title_style))
        story.append(Spacer(1, 12))
        
        # Metadata
        story.append(Paragraph(f"<b>Original File:</b> {questionnaire.original_filename}", styles['Normal']))
        story.append(Paragraph(f"<b>Export Date:</b> {datetime.now().strftime('%Y-%m-%d %H:%M')}", styles['Normal']))
        story.append(Paragraph(f"<b>Status:</b> {questionnaire.status}", styles['Normal']))
        story.append(Spacer(1, 20))
        
        # Questions and Answers
        questions = Question.query.filter_by(questionnaire_id=questionnaire.id).order_by(Question.question_number).all()
        
        for question in questions:
            # Question
            q_text = f"<b>Q{question.question_number}:</b> {self._escape_xml(question.question_text)}"
            story.append(Paragraph(q_text, question_style))
            
            # Answer
            if question.answers:
                answer = question.answers[0]
                a_text = f"<b>Answer:</b> {self._escape_xml(answer.answer_text)}"
                story.append(Paragraph(a_text, answer_style))
                
                # Citations
                if answer.citations:
                    citation_texts = []
                    for citation in answer.citations:
                        cite_str = f"{citation.document.original_filename}"
                        if citation.page_number:
                            cite_str += f", p.{citation.page_number}"
                        citation_texts.append(cite_str)
                    c_text = f"<i>Citations: {'; '.join(citation_texts)}</i>"
                    story.append(Paragraph(c_text, citation_style))
                
                # Confidence
                if answer.confidence_score is not None:
                    conf_text = f"<i>Confidence: {answer.confidence_score:.0%}</i>"
                    story.append(Paragraph(conf_text, citation_style))
                
                # Edited indicator
                if answer.is_edited:
                    story.append(Paragraph("<i>(Edited by user)</i>", citation_style))
            else:
                story.append(Paragraph("<b>Answer:</b> Not answered", answer_style))
        
        doc.build(story)
        
        return output_path
    
    def _escape_xml(self, text: str) -> str:
        """Escape XML special characters."""
        return text.replace('&', '&amp;').replace('<', '&lt;').replace('>', '&gt;').replace('"', '&quot;')
